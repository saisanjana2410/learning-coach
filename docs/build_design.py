from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
d=Document()
s=d.sections[0]; s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Heading 1'].font.size=Pt(19)
d.styles['Heading 2'].font.size=Pt(12)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def bullets(items):
 for t in items: d.add_paragraph(t,'List Bullet')
def page(title):
 if len(d.paragraphs)>0: d.add_page_break()
 d.add_heading(title,1)
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
 for i,x in enumerate(headers): t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row): cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   if widths: c.width=Inches(widths[ci])
   pr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F3F6F8' if ri%2==0 else 'FFFFFF')); pr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5); para.paragraph_format.space_before=Pt(5)
    for r in para.runs: r.font.size=Pt(9); r.bold=ri==0
  if ri==0:
   repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
  keep=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(keep)
 d.add_paragraph()

d.add_heading('Adaptive Learning and Interview Preparation Agent',0)
d.add_paragraph('Architecture and implementation design',style='Subtitle')
p('Version 1.0 | October 22 2026 | Design baseline')
p('This document defines a personal learning coach that teaches from basics to advanced, supports applied and interview practice, remembers results, identifies specific weak areas, and schedules revision. It is the implementation reference for a Python and LangChain backend with a basic React frontend.')
p('The first release delivers the complete learning loop for linear regression. The architecture supports a broader AI and ML curriculum without requiring that curriculum to be built before the first usable release.')
h('Core design decision')
p('Use one tutor agent and a separate structured answer evaluation call. Keep storage, session transitions, progression rules, and review scheduling in deterministic Python services. Build the frontend and backend against the same API contracts.')
h('Reading guide')
table(['Section','Coverage'],[
('1','Product scope and requirements'),('2','Architecture and technology'),('3','User flows and frontend'),('4','Agent behavior and tools'),('5','Data model and relationships'),('6','Assessment and revision rules'),('7','API contracts and execution'),('8','Reliability and evaluation'),('9','Delivery plan and future scope'),('10','Decision register and references')],[.7,5.9])
p('Audience: project owner and developer. The design assumes a local, single-learner application. Proposed policies and undecided choices are identified explicitly; this document does not represent implemented behavior.')

page('1 Product scope and requirements')
h('Problem and intended outcome')
p('Learners often have fragmented notes, inconsistent practice, and little evidence of what they can explain or apply. The coach connects instruction, exercises, interview answers, and revision through a shared skill history. Its goal is to guide deliberate practice and expose progress through evidence rather than claim an interview readiness score.')
h('First release')
bullets(['Onboard one learner with target role, experience, timezone, and study availability.','Provide Learn, Practice, and Interview modes for linear regression.','Assess prerequisites and teach small lessons with examples and follow-up questions.','Accept written explanations and code as text; give rubric-based feedback.','Persist attempts and dimension-specific gaps across sessions.','Show due revisions, recent activity, and suggested next steps in a basic web interface.','Allow the learner to override recommendations and resume an interrupted session.'])
h('Initial curriculum')
p('Cover predictions and coefficients, slope and intercept, residuals, mean squared error, train and validation splits, basic evaluation, and assumptions. Include prerequisite checks on basic algebra and Python. Introduce gradient descent conceptually after the essential prediction and error concepts are demonstrated.')
h('Boundaries')
p('The initial application does not execute submitted code, browse for news, ingest arbitrary documents, send notifications, or connect to job platforms. Code feedback is a review of text, not proof that the program runs. Advanced curriculum coverage is an expansion goal.')
h('Functional acceptance')
table(['Requirement','Observable outcome'],[
('Persistent learning','An attempt and its feedback remain after restart.'),('Shared weak areas','An interview gap appears in practice recommendations.'),('Adaptive instruction','A misconception triggers a targeted explanation and fresh exercise.'),('Revision','A due review appears with its scheduling reason.'),('Learner control','A suggested topic can be skipped without erasing its history.'),('Traceable feedback','Feedback references a rubric criterion and answer evidence.')],[1.5,5.1])

page('2 Architecture and technology')
p('React communicates with FastAPI over JSON. FastAPI invokes the session controller, which coordinates content retrieval, the tutor, evaluation, and learning services. SQLite is the authoritative record of learning progress.')
table(['Layer','Choice','Responsibility'],[
('Frontend','React with Vite and plain CSS','Dashboard, learning workspace, progress and onboarding'),('API','FastAPI','Validate requests, resolve learner context, expose JSON contracts'),('Agent','LangChain create_agent','Tutor conversation and bounded tool selection'),('Model','One hosted tool-capable LLM','Teaching and structured answer evaluation'),('Validation','Pydantic','Input and output schemas'),('Persistence','SQLite with SQLAlchemy','Learning records and session state'),('Content','Markdown and JSON','Versioned lessons, questions and grading rubrics'),('Tests','pytest and reviewed examples','Service correctness and feedback evaluation')],[1,2,3.6])
h('Request path')
p('Browser → FastAPI → session controller → tutor or evaluator → learning service → database. Read-only agent tools access the content catalog and learner-scoped history through services. The browser never accesses the LLM provider or database directly.')
h('Agent boundary')
p('The tutor chooses explanations, appropriate content, and follow-up questions. The controller binds the current question and mode, caps tool steps, and controls state transitions. A structured evaluation call assesses submitted answers using the stored question and rubric. A validated service transaction saves results and computes progress and review changes.')
h('Why this is an agent')
p('Within a session, the tutor can inspect past attempts, retrieve a lesson, choose a suitable question, and adapt its next action to the learner response. Required bookkeeping is application behavior and does not depend on discretionary tool calls.')
h('Deployment and configuration')
p('Develop locally with separate frontend and API processes. Keep secrets in backend environment variables and exclude them from source control. Restrict CORS to the configured frontend origin. Bind the initial single-user backend to localhost. Model provider, model name, token limits, timeouts, and database path are configuration values. Pin compatible package versions during implementation.')

page('3 User flows and frontend')
h('First visit')
p('Onboarding → choose target role and study availability → select linear regression → short diagnostic → see assessed and unassessed skills → begin a lesson or choose practice. The learner may skip the diagnostic; skipped skills remain unassessed.')
h('Learning session')
p('Dashboard → choose a due review or next topic → select mode → receive content → ask questions or submit an answer → receive feedback → complete a fresh follow-up → view summary and revision queue. Asking for clarification is a chat message; submitting an assessed answer is a distinct action.')
h('Return visit')
p('Load profile and durable history → show due reviews and active session → resume or start a session → reassess selected gaps → adjust review schedule. An overdue topic remains due until reviewed or explicitly deferred; missed days do not create duplicate reviews.')
table(['View','Main elements','States to support'],[
('Onboarding','Role, experience, daily minutes, timezone','New profile, validation error, saved'),('Dashboard','Continue session, due reviews, weak areas, suggested topic','New learner, loading, empty queue, error'),('Workspace','Mode, topic, lesson or question, chat, answer box, feedback','Thinking, awaiting answer, hints, evaluation pending, retry'),('Progress','Skill dimensions, evidence, attempt history, next review','Unassessed, learning, needs review, proficiency')],[1.1,3,2.5])
h('Mode rules')
p('Learn mode permits explanations, worked examples, and guided exercises. Practice mode offers questions with optional hints. Interview mode defers coaching until answer submission; choosing to reveal a hint records assistance. Interview communication is assessed from written answers only in the first release.')
h('Interface behavior')
p('Use a desktop layout with simple navigation and a central conversation area; stack panels on narrow screens. Show feedback as correct points, gaps, evidence, and next action. Use status labels and text as well as color. Preserve unsent answer text after recoverable network errors. Disable duplicate submission while a request is pending.')
h('Session states')
p('ACTIVE → AWAITING_ANSWER → EVALUATING → FEEDBACK_READY → AWAITING_ANSWER or COMPLETED. A model failure produces a retryable error without discarding the submitted attempt. The controller owns transitions; natural language from the model cannot directly change them.')

page('4 Agent behavior and tools')
h('Tutor instructions')
p('Teach one manageable concept at a time, retrieve relevant reference material, and adapt detail to the learner. Distinguish missing knowledge from an incorrect belief. Use history to avoid repeating the same question as evidence of improvement. Explain why a review is recommended. If evidence is insufficient, say so and gather another answer.')
table(['Tool','Inputs','Output'],[
('get_learning_profile','Learner identity from server context','Goals, skill evidence and recent activity'),('get_due_reviews','Optional topic filter and limit','Due skills, dates and scheduling reasons'),('get_lesson','Skill ID and optional lesson ID','Reviewed lesson text and content version'),('get_question','Skill ID, mode, difficulty','Question ID, prompt and permitted hints'),('get_attempt_history','Skill ID and bounded limit','Prior answers, gap summaries and assistance flags')],[1.65,2,2.95])
p('All tools are read-only and learner-scoped. The server supplies learner identity; it is not a model-controlled argument. Validate identifiers, constrain result sizes, and reject unknown skills. Return structured error results when content is unavailable.')
h('Protected assessment context')
p('The question service stores the selected question ID on the session. Learner-facing tools omit reference answers and rubrics. The evaluator receives the server-loaded rubric after submission. This reduces accidental disclosure but cannot guarantee that a generative tutor never reveals an answer; mode-specific behavior must be evaluated.')
h('Evaluation output')
p('Require a validated object containing question ID, overall result, per-skill and per-dimension judgments, rubric criterion outcomes, answer evidence, misconceptions, strengths, improvement guidance, and suggested follow-up skill. Add evaluator model and prompt versions on the server. Treat uncertainty as an explicit insufficient-evidence result.')
h('Write operations')
p('The controller saves the attempt before calling the evaluator. After valid evaluation, the learning service commits assessment, skill evidence, and review changes in one transaction. The model suggests learning actions but cannot issue arbitrary SQL, update dates, or mark proficiency directly.')
h('Execution limits')
p('Apply bounded tool steps, context size, request duration, and output tokens. Load a short session summary and relevant history rather than the entire lifetime conversation. Record model usage and tool outcomes without storing secrets. No external browsing tools are exposed in release one.')

page('5 Data model and relationships')
p('Use UUID identifiers, foreign keys, UTC timestamps, and schema migrations. Display dates in the learner timezone. Keep content versions immutable once attempts reference them. SQLAlchemy associations implement many-to-many relationships between questions and skills and between lessons and skills.')
table(['Entity','Important fields'],[
('Learner','id, target_role, experience, study_minutes, timezone, created_at'),('Topic and Skill','topic_id; skill_id, topic_id, name; prerequisite edges'),('Lesson','id, version, content_path; lesson_skill association'),('Question','id, version, mode, difficulty, prompt, rubric, hints; question_skill association'),('Session','id, learner_id, mode, state, current_question_id, started_at, ended_at'),('Message','id, session_id, role, content, sequence, created_at'),('Attempt','id, session_id, question_id, question_version, answer, hints_used, status, idempotency_key'),('Assessment','id, attempt_id, result, feedback, model_version, prompt_version, rubric_version'),('SkillEvidence','assessment_id, skill_id, dimension, criterion_result, evidence_excerpt, assistance'),('SkillProgress','learner_id, skill_id, dimension, status, evidence_count, last_assessed_at, rule_version'),('ReviewItem','learner_id, skill_id, dimension, due_at, interval_days, reason, scheduling_version'),('ModelRun','id, session_id, attempt_id optional, purpose, status, latency, token_usage, error_code')],[1.35,5.25])
h('Relationships and constraints')
p('Learner has many sessions; session has messages and attempts; attempt references one exact question version. Each attempt has at most one accepted assessment in the first release. Assessment has multiple skill evidence records. SkillProgress and ReviewItem each have a unique learner, skill, dimension key. A skill can have multiple prerequisite skills; reject prerequisite cycles during content validation.')
h('Evidence and derived state')
p('Attempts and accepted assessments are retained as source evidence. SkillProgress is a derived summary and can be recalculated under a new rule version. ReviewItem represents the current scheduled review, not a duplicate event for every missed day. New assessments retain the prior attempt history.')

page('6 Assessment and revision rules')
h('Track dimensions separately')
p('Track conceptual understanding, practical application, coding, and interview communication independently. A coding attempt may assess conceptual knowledge too, but only where the rubric explicitly supports it. Written communication feedback is not evidence about spoken delivery.')
h('Rubric and confidence')
p('Each criterion is marked met, partially met, not met, or insufficient evidence, with a supporting answer excerpt. Questions specify which skills and dimensions they assess. Model confidence is not a calibrated probability of learner mastery. Uncertain evaluations request clarification or another question and do not trigger negative progression updates.')
h('Initial progression policy')
p('Not assessed means there is no usable evidence. Learning means evidence exists but proficiency has not been demonstrated. Needs review means a clear gap was observed. Demonstrated proficiency initially requires two independent successful answers on different questions across separate sessions, including one after a delay of at least four days. A later clear error returns the affected dimension to needs review. These thresholds are proposed rules to test and adjust.')
table(['Result','Immediate behavior','Next review'],[
('Incorrect or major gap','Explain gap and offer guided retry','1 day'),('Partial or heavily assisted','Give targeted practice','2 days'),('Correct independently','Offer next suitable skill','4 days'),('Fresh delayed independent success','Extend existing successful interval','8 days then 16 days'),('Insufficient evidence','Ask clarification or new question','Keep prior schedule')],[1.65,2.8,2.15])
h('Scheduling details')
p('Compute dates in the learning service. A guided retry in the same session does not count as delayed retention evidence. Repeated exposure to the same question does not satisfy the fresh-question rule. When a new gap appears, shorten the relevant review interval. Merge conceptual and interview review recommendations for the same skill in the UI while keeping their evidence distinct.')
h('Priority and learner control')
p('Prioritize overdue reviews, then recurring gaps, then prerequisites for the selected topic. Limit the suggested workload to the session time budget. A learner may defer a review or advance manually; show that decision without manufacturing proficiency. Explain each due date using its result and scheduling rule.')

page('7 API contracts and execution')
p('Use JSON under /api/v1. Resolve the local learner on the server. A future authenticated release will replace that resolver without accepting arbitrary learner IDs from client requests. Initial model calls are request and response; token streaming can be added later.')
table(['Method and route','Input','Output'],[
('GET /profile','None','Learner settings'),('PUT /profile','Role, experience, time, timezone','Validated saved profile'),('GET /dashboard','None','Active session, due reviews, weak areas'),('GET /topics','None','Available modules and skill coverage'),('POST /sessions','Mode, topic or skill, time budget','Session and first activity'),('GET /sessions/{id}','Session identifier','State, messages and active activity'),('POST /sessions/{id}/messages','Client message ID, text','Assistant message and activity'),('POST /sessions/{id}/attempts','Question ID, answer, idempotency key','Attempt status, feedback and review changes'),('GET /attempts/{id}','Attempt identifier','Pending, failed or accepted assessment'),('POST /sessions/{id}/complete','None','Session summary'),('GET /progress','Optional topic filter','Skill dimensions and attempt summaries'),('POST /reviews/{id}/defer','New date and optional reason','Updated review')],[2.25,1.8,2.55])
h('Submission contract')
p('Client sends the active question ID and a unique idempotency key. The server checks that the question matches the session and is awaiting an answer, then saves one attempt. It evaluates against the stored rubric, validates the response, and commits results. Repeating the same key returns the existing result. Reusing a key with different content is a conflict.')
h('Failures and retries')
p('Validation failures return 422; an incompatible state returns 409; missing records return 404. Provider timeouts or transient failures return a stable error code, request ID, and retryability flag. A pending attempt can be retrieved by ID. Retrying evaluation reuses the saved answer and cannot create duplicate accepted evidence. Do not mark a failed evaluation as an incorrect answer.')

page('8 Reliability and evaluation')
h('Persistence and session recovery')
p('Use a database transaction for accepted assessment and derived updates. Store user messages before model work and associate responses with client message IDs. Serialize changes within a session to avoid overlapping submissions. On restart, reconstruct the active activity from the database; transient conversation context is not the learning record.')
h('Content and privacy')
p('Validate curriculum references, prerequisite graphs, question coverage, and version IDs before import. Send only the relevant history to the model provider. Keep provider keys on the backend and omit answer text from routine operational logs. The initial app stores learning data locally; selected content and answers still leave the device when sent to the hosted model. Add export and delete operations before broader user distribution.')
h('Meaningful tests')
table(['Test area','Evidence required'],[
('Learning services','History survives restart; per-dimension updates do not overwrite unrelated skills.'),('Review rules','Due dates, delayed evidence, assistance and insufficient evidence follow policy.'),('API behavior','Invalid state rejected; duplicate requests do not duplicate results.'),('Agent tools','Learner scope enforced; question tools exclude protected rubric fields.'),('Feedback quality','Compare criterion judgments with human-reviewed sample answers.'),('End to end','Interview gap appears in dashboard and leads to a fresh practice review.')],[1.4,5.2])
h('Evaluation set')
p('Begin with at least 20 reviewed answers spanning correct, partial, incorrect, ambiguous, and hint-assisted cases. Record criterion agreement with the reference judgments, unsupported-feedback frequency, and validity of follow-up skill selection. Track latency and token usage. Run baseline and revised evaluator prompts on the same set. Report sample size and limitations alongside results.')
h('Completion criteria')
p('The first release is complete when the end-to-end learning loop works through the web interface, data survives restart, failure recovery and duplicate submission checks pass, and reviewed evaluation results are documented. Establish a measured quality baseline before selecting a numerical feedback accuracy target. Do not invent resume performance numbers.')

page('9 Delivery plan and future scope')
table(['Milestone','Backend and content','Frontend and verification'],[
('1 Foundation','Project structure, profile, schema and API contracts','Navigation, onboarding, mock API fixtures'),('2 First lesson','Versioned linear regression content and read tools','Dashboard and learning workspace'),('3 Tutor loop','LangChain tutor, messages and session recovery','Conversation, loading and retry states'),('4 Practice feedback','Rubrics, evaluator, attempt persistence','Answer submission and feedback display'),('5 Adaptive revision','Skill evidence, progression and review scheduler','Weak areas, review queue and progress history'),('6 Interview mode','No-coaching rules and shared history','Mode selector and session summary'),('7 Release validation','Evaluation set, integration checks and configuration guide','End-to-end checks and basic responsive layout')],[1.1,2.75,2.75])
h('Working in parallel')
p('Define request and response schemas first. Frontend development uses sample responses with the same shapes, including empty and error states. Replace each mock endpoint with its real API as the corresponding service becomes available. Keep the same user flow throughout integration rather than building a separate demonstration interface.')
h('Suggested project structure')
p('frontend/src contains pages, components, api and types. backend/app contains api, schemas, agent, services, models and configuration. content contains lessons, questions and rubrics. tests contains service, API and evaluation fixtures. docs contains design and setup material. Keep prompts versioned alongside backend agent configuration.')
h('Expansion order')
bullets(['Expand the curriculum and measure performance across more topics.','Add job-description skill mapping and targeted interview sessions.','Execute coding exercises in an isolated runner with resource limits.','Add sourced technology updates with publication dates and relevance explanations.','Add voice, proactive reminders, authentication and deployment when the core workflow is reliable.'])
p('Automatic reminders require a real scheduler and notification channel. The first version only stores due reviews and shows them when the learner opens the application. Technology news requires source retrieval and verification; it should not rely on model memory.')

page('10 Decision register and references')
table(['Decision','Status','Rationale or next action'],[
('Single tutor plus evaluator','Design baseline','Small, understandable agent architecture'),('React and FastAPI','Agreed direction','Build a basic web interface alongside backend'),('SQLite and SQLAlchemy','Design baseline','Persistent local data with clear relationships'),('Linear regression first','Design baseline','One complete module before curriculum expansion'),('Model provider and model','Open','Choose using available account, cost and evaluation results'),('Question bank content','To author','Review lesson accuracy and rubric coverage'),('Revision thresholds','Proposed policy','Evaluate usability and adjust from evidence'),('Local application only','First-release boundary','Authentication required before network deployment'),('Streaming responses','Deferred','Request and response is sufficient initially'),('Product name','Open','Use the descriptive working title for now')],[1.9,1.3,3.4])
h('Architecture tradeoffs')
p('A small curated catalog makes retrieval predictable but limits coverage. A hosted LLM reduces setup effort but introduces cost, latency, and external processing. Rule-based revision is explainable but only a starting heuristic. SQLite simplifies local use; shared deployment may require a server database and stronger concurrency controls. Text-based code review enables early practice but cannot validate runtime correctness.')
h('Official implementation references')
p('LangChain create_agent reference: https://reference.langchain.com/python/langchain/agents/factory/create_agent')
p('LangChain structured output documentation: https://github.com/langchain-ai/docs/blob/main/src/oss/langchain/structured-output.mdx')
p('LangChain short term memory documentation: https://github.com/langchain-ai/docs/blob/main/src/oss/langchain/short-term-memory.mdx')
p('These references support the framework interface and memory concepts discussed in the design. Scheduling thresholds, application schemas, user flows, and quality gates are project design decisions rather than claims made by those sources. Verify package compatibility when dependencies are pinned.')
h('Next implementation step')
p('Choose the model provider, then define API schemas and a small reviewed content pack. Begin the frontend shell and backend profile service together. Keep the learning loop usable at each milestone and update this design when implementation decisions change.')

footer=s.footer.paragraphs[0]; footer.alignment=2
r=footer.add_run('Page '); r.font.size=Pt(9)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
d.core_properties.title='Adaptive Learning and Interview Preparation Agent Architecture'
d.core_properties.subject='Product scope architecture user flows tools and data model'
d.core_properties.author=''
out=ROOT/'Learning_Agent_Architecture.docx'; d.save(out)
print(out)
